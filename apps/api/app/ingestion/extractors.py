"""Extractors: VTT/SRT, PDF, PPTX, plain text/notes/code."""

from __future__ import annotations

import re
from pathlib import Path

from app.ingestion.models import Segment

VTT_TIMING = re.compile(
    r"(?P<h1>\d{2}):(?P<m1>\d{2}):(?P<s1>\d{2})[.,](?P<ms1>\d{3})\s*-->\s*"
    r"(?P<h2>\d{2}):(?P<m2>\d{2}):(?P<s2>\d{2})[.,](?P<ms2>\d{3})"
)


def _to_ms(h: str, m: str, s: str, ms: str) -> int:
    return ((int(h) * 3600 + int(m) * 60 + int(s)) * 1000) + int(ms)


def parse_caption_file(path: str | Path, source_type: str = "video", video_id: str | None = None) -> list[Segment]:
    """Parse VTT or SRT into timestamped segments."""
    p = Path(path)
    try:
        import webvtt  # type: ignore

        segments: list[Segment] = []
        for i, caption in enumerate(webvtt.read(str(p))):
            start = _webvtt_ts_to_ms(caption.start)
            end = _webvtt_ts_to_ms(caption.end)
            text = _strip_tags(caption.text)
            if text.strip():
                segments.append(
                    Segment(
                        text=text.strip(),
                        source_type=source_type,
                        video_id=video_id,
                        start_ms=start,
                        end_ms=end,
                        index=i,
                    )
                )
        if segments:
            return segments
    except (ImportError, OSError, ValueError, AttributeError):
        pass
    return _parse_caption_stdlib(p, source_type, video_id)


def _webvtt_ts_to_ms(ts: str) -> int:
    ts = ts.strip().replace(",", ".")
    parts = ts.split(":")
    if len(parts) == 3:
        h, m, rest = parts
    else:
        h, (m, rest) = "0", parts
    s, _, ms = rest.partition(".")
    return _to_ms(h, m, s, (ms + "000")[:3])


def _strip_tags(text: str) -> str:
    return re.sub(r"<[^>]+>", "", text).replace("\n", " ")


def _parse_caption_stdlib(p: Path, source_type: str, video_id: str | None) -> list[Segment]:
    raw = p.read_bytes()
    try:
        import chardet

        enc = chardet.detect(raw).get("encoding") or "utf-8"
    except (ImportError, AttributeError):
        enc = "utf-8"
    text = raw.decode(enc, errors="replace")
    segments: list[Segment] = []
    idx = 0
    for block in re.split(r"\n\s*\n", text):
        lines = [ln.strip() for ln in block.strip().splitlines() if ln.strip()]
        if not lines:
            continue
        timing_i = next((i for i, ln in enumerate(lines) if VTT_TIMING.search(ln)), None)
        if timing_i is None:
            continue
        m = VTT_TIMING.search(lines[timing_i])
        assert m is not None
        start = _to_ms(m["h1"], m["m1"], m["s1"], m["ms1"])
        end = _to_ms(m["h2"], m["m2"], m["s2"], m["ms2"])
        body = " ".join(lines[timing_i + 1 :])
        body = _strip_tags(body)
        if body.strip():
            segments.append(
                Segment(
                    text=body.strip(),
                    source_type=source_type,
                    video_id=video_id,
                    start_ms=start,
                    end_ms=end,
                    index=idx,
                )
            )
            idx += 1
    return segments


def extract_pdf(path: str | Path, source_type: str = "pdf") -> list[Segment]:
    try:
        import fitz  # pymupdf
    except ImportError:
        import pymupdf as fitz  # type: ignore

    segments: list[Segment] = []
    with fitz.open(str(path)) as doc:
        for page_index in range(len(doc)):
            text = doc[page_index].get_text().strip()
            if text:
                segments.append(
                    Segment(
                        text=text,
                        source_type=source_type,
                        page_no=page_index + 1,
                        index=page_index,
                    )
                )
    return segments


def extract_pptx(path: str | Path, source_type: str = "pptx") -> list[Segment]:
    from pptx import Presentation

    segments: list[Segment] = []
    prs = Presentation(str(path))
    for i, slide in enumerate(prs.slides):
        parts: list[str] = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                t = shape.text_frame.text.strip()
                if t:
                    parts.append(t)
        text = "\n".join(parts).strip()
        if text:
            segments.append(
                Segment(
                    text=text,
                    source_type=source_type,
                    slide_no=i + 1,
                    index=i,
                )
            )
    return segments


def extract_text(path: str | Path, source_type: str = "note") -> list[Segment]:
    raw = Path(path).read_bytes()
    try:
        import chardet

        enc = chardet.detect(raw).get("encoding") or "utf-8"
    except ImportError:
        enc = "utf-8"
    text = raw.decode(enc, errors="replace")
    paragraphs = [p.strip() for p in re.split(r"\n\s*\n", text) if p.strip()]
    segments: list[Segment] = []
    for i, para in enumerate(paragraphs):
        segments.append(Segment(text=para, source_type=source_type, index=i))
    if not segments and text.strip():
        segments.append(Segment(text=text.strip(), source_type=source_type, index=0))
    return segments


def extract(path: str | Path, resource_type: str, video_id: str | None = None) -> list[Segment]:
    p = Path(path)
    suffix = p.suffix.lower()
    if resource_type == "video" or suffix in (".vtt", ".srt"):
        return parse_caption_file(p, source_type=resource_type, video_id=video_id)
    if resource_type == "pdf" or suffix == ".pdf":
        return extract_pdf(p)
    if resource_type == "pptx" or suffix in (".pptx", ".ppt"):
        return extract_pptx(p)
    return extract_text(p, source_type=resource_type)
