"""Document extraction for pdf/pptx.

Documents are NOT time-gated (``app/retrieval``), so segments here carry
``page_no``/``slide_no`` and ``start_ms``/``end_ms`` of ``None`` -- there is no
moment-in-time bound to enforce. Keeping the fields ``None`` rather than 0 is
what makes "docs are ungated" a structural fact rather than a special case in
the filter.
"""

from __future__ import annotations

import logging
from dataclasses import dataclass, field
from pathlib import Path

from app.ingestion import ocr as ocr_mod
from app.ingestion.chunking import TextSegment

log = logging.getLogger(__name__)


@dataclass
class DocumentPassReport:
    """What a document ingest actually managed to read.

    The same honesty contract as :class:`~app.ingestion.visual.VisualPassReport`:
    a PDF with no text layer used to produce ``[]`` and therefore a clean
    ``status: done`` with zero chunks -- indistinguishable from a document that
    was read successfully and happened to be short. This report is what makes
    the two tell themselves apart, and it is carried onto the job payload.
    """

    segments: list[TextSegment] = field(default_factory=list)
    #: The OCR pass, skipped unless the text layer came back (near) empty.
    ocr: ocr_mod.OcrPassReport = field(default_factory=ocr_mod.OcrPassReport)

    @property
    def text_chars(self) -> int:
        return sum(len(s.text) for s in self.segments)

    @property
    def status(self) -> str:
        """``complete`` | ``partial`` | ``failed``, as in the visual pass."""
        if not self.segments:
            return "failed"
        if self.ocr.attempted and self.ocr.status != "complete":
            return self.ocr.status
        return "complete"

    @property
    def warnings(self) -> list[str]:
        return list(self.ocr.warnings)


def extract_pdf_pass(path: str | Path) -> DocumentPassReport:
    """Extract a PDF, falling back to OCR only when there is no text layer.

    A PDF that already has a text layer is never sent to the model: OCR is
    slow, expensive, and strictly worse than the text layer it would replace.
    Only when ordinary extraction yields fewer than
    :data:`app.ingestion.ocr.OCR_MIN_TEXT_CHARS` characters is OCR considered,
    and then only if an operator has opted in.
    """
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    out: list[TextSegment] = []
    for index, page in enumerate(reader.pages, start=1):
        try:
            text = (page.extract_text() or "").strip()
        except Exception as exc:  # noqa: BLE001 - one bad page must not fail all
            log.warning("pdf page %d of %s unreadable: %s", index, path, exc)
            continue
        if text:
            out.append(TextSegment(text=text, page_no=index, kind="document"))
    page_count = len(reader.pages)

    if sum(len(s.text) for s in out) >= ocr_mod.OCR_MIN_TEXT_CHARS:
        # Text layer present. Not attempted, and that is not a failure.
        log.info("extracted %d page segment(s) from %s", len(out), path)
        return DocumentPassReport(segments=out)

    ocr_report = ocr_mod.extract_ocr_pass(path, pages_expected=page_count)
    if ocr_report.result is not None and ocr_report.result.ok:
        result = ocr_report.result
        if result.pages:
            out = [
                TextSegment(text=page_text, page_no=number, kind="document")
                for number, page_text in enumerate(result.pages, start=1)
                if page_text.strip()
            ]
        elif result.text.strip():
            # The backend returned one merged document rather than per-page
            # output. Stored without a page number rather than attributed to a
            # page we cannot prove it came from.
            out = [TextSegment(text=result.text.strip(), kind="document")]
        log.info(
            "no text layer in %s; OCR supplied %d page segment(s)",
            path, len(out),
        )
        return DocumentPassReport(segments=out, ocr=ocr_report)

    log.error(
        "no extractable text in %s and no OCR result (status=%s): %s",
        path, ocr_report.status, ocr_report.reason or "no reason recorded",
    )
    return DocumentPassReport(segments=out, ocr=ocr_report)


def extract_pdf(path: str | Path) -> list[TextSegment]:
    """Segments only. Thin wrapper over :func:`extract_pdf_pass`.

    Prefer the report form in the ingest pipeline -- it is the only way to see
    whether the document was actually read.
    """
    return extract_pdf_pass(path).segments


def extract_pptx(path: str | Path) -> list[TextSegment]:
    from pptx import Presentation

    deck = Presentation(str(path))
    out: list[TextSegment] = []
    for index, slide in enumerate(deck.slides, start=1):
        parts: list[str] = []
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                text = (shape.text_frame.text or "").strip()
                if text:
                    parts.append(text)
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    cells = [c.text.strip() for c in row.cells if c.text.strip()]
                    if cells:
                        parts.append(" | ".join(cells))
        if parts:
            out.append(
                TextSegment(text="\n".join(parts), slide_no=index, kind="document")
            )
    log.info("extracted %d slide segment(s) from %s", len(out), path)
    return out


def extract(path: str | Path, source_type: str) -> list[TextSegment]:
    """Dispatch on the declared source type. Raises on an unknown type."""
    path = Path(path)
    if source_type == "pdf":
        return extract_pdf(path)
    if source_type == "pptx":
        return extract_pptx(path)
    raise ValueError(f"extract() handles documents only, got {source_type!r}")